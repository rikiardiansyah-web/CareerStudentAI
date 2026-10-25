import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os

TEMPLATE_PATH = r'C:\Users\Admin\.gemini\antigravity\brain\8c1f2961-a2f7-42ef-86e0-6c72fe3fc1a0\.user_uploaded\media_1792948346046_0a1e617a.docx'
LOGO_PATH = r'C:\Users\Admin\.gemini\antigravity\scratch\careerstudent-ai\public\icons\logo-cs.jpg'
SCREENSHOT1_PATH = r'C:\Users\Admin\.gemini\antigravity\brain\8c1f2961-a2f7-42ef-86e0-6c72fe3fc1a0\.user_uploaded\media_1792359115071_ae73705e.jpg'
SCREENSHOT2_PATH = r'C:\Users\Admin\.gemini\antigravity\brain\8c1f2961-a2f7-42ef-86e0-6c72fe3fc1a0\.user_uploaded\media_1792361157177_53c3abec.jpg'

OUTPUT_SCRATCH = r'C:\Users\Admin\.gemini\antigravity\scratch\careerstudent-ai\Deskripsi_Produk_CareerStudentAI.docx'
OUTPUT_BRAIN = r'C:\Users\Admin\.gemini\antigravity\brain\8c1f2961-a2f7-42ef-86e0-6c72fe3fc1a0\Deskripsi_Produk_CareerStudentAI.docx'

doc = docx.Document(TEMPLATE_PATH)

# ==========================================
# 1. PARAGRAPH 0: Header
# ==========================================
# Keeps "KARYA 01  /  LOMBA INOVASI DIGITAL"

# ==========================================
# 2. PARAGRAPH 1: Product Title
# ==========================================
p1 = doc.paragraphs[1]
p1.text = ""
p1.paragraph_format.space_before = Pt(6)
p1.paragraph_format.space_after = Pt(4)
p1.paragraph_format.line_spacing = 1.15
r1 = p1.add_run("CAREERSTUDENT AI: ADAPTIVE CAREER CO-PILOT BERBASIS HUMAN-IN-THE-LOOP & SOCRATIC DIALOGUE UNTUK MENJEMBATANI KESIAPAN KERJA MAHASISWA MULTIDISIPLIN DI ERA AI GENERATIF")
r1.font.name = "Calibri"
r1.font.size = Pt(15)
r1.bold = True
r1.font.color.rgb = RGBColor(0x17, 0x32, 0x4D)

# ==========================================
# 3. PARAGRAPH 2: Team Name
# ==========================================
p2 = doc.paragraphs[2]
p2.text = ""
p2.paragraph_format.space_before = Pt(2)
p2.paragraph_format.space_after = Pt(8)
r2_label = p2.add_run("NAMA TIM : ")
r2_label.font.name = "Calibri"
r2_label.font.size = Pt(11)
r2_label.bold = True
r2_label.font.color.rgb = RGBColor(0x17, 0x32, 0x4D)

r2_val = p2.add_run("TIM AQUAMAN TEKNOKRAT")
r2_val.font.name = "Calibri"
r2_val.font.size = Pt(11)
r2_val.bold = True
r2_val.font.color.rgb = RGBColor(0x00, 0xA6, 0xA6)

# ==========================================
# 4. TABLE 0: Logo Inovasi
# ==========================================
t0 = doc.tables[0]
c0 = t0.rows[0].cells[0]
p_logo = c0.paragraphs[0]
p_logo.text = ""
p_logo.paragraph_format.space_before = Pt(4)
p_logo.paragraph_format.space_after = Pt(4)
p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_logo = p_logo.add_run()
r_logo.add_picture(LOGO_PATH, width=Inches(1.8))

p_logo_label = c0.add_paragraph()
p_logo_label.paragraph_format.space_before = Pt(2)
p_logo_label.paragraph_format.space_after = Pt(4)
p_logo_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_logo_lbl = p_logo_label.add_run("Logo Resmi Inovasi CareerStudent AI")
r_logo_lbl.font.name = "Calibri"
r_logo_lbl.font.size = Pt(8.5)
r_logo_lbl.font.color.rgb = RGBColor(0x17, 0x32, 0x4D)
r_logo_lbl.bold = True

# ==========================================
# 5. TABLE 1: Screenshots Produk
# ==========================================
t1 = doc.tables[1]
c1 = t1.rows[0].cells[0]
p_ss1 = c1.paragraphs[0]
p_ss1.text = ""
p_ss1.paragraph_format.space_before = Pt(6)
p_ss1.paragraph_format.space_after = Pt(3)
p_ss1.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_ss1 = p_ss1.add_run()
r_ss1.add_picture(SCREENSHOT1_PATH, width=Inches(5.3))

p_cap1 = c1.add_paragraph()
p_cap1.paragraph_format.space_before = Pt(2)
p_cap1.paragraph_format.space_after = Pt(8)
p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_cap1 = p_cap1.add_run("Gambar 1. Antarmuka Modul ATS-Compliant Fast CV Builder dengan Explicit Approval Gate & Socratic Prompting")
r_cap1.font.name = "Calibri"
r_cap1.font.size = Pt(8.5)
r_cap1.italic = True
r_cap1.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

p_ss2 = c1.add_paragraph()
p_ss2.paragraph_format.space_before = Pt(4)
p_ss2.paragraph_format.space_after = Pt(3)
p_ss2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_ss2 = p_ss2.add_run()
r_ss2.add_picture(SCREENSHOT2_PATH, width=Inches(5.3))

p_cap2 = c1.add_paragraph()
p_cap2.paragraph_format.space_before = Pt(2)
p_cap2.paragraph_format.space_after = Pt(6)
p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_cap2 = p_cap2.add_run("Gambar 2. Modul Analisis Skill-Gap Matrix (Taksonomi WEF 2023) & Sesi Interactive Mock Interview Berbasis STAR Framework")
r_cap2.font.name = "Calibri"
r_cap2.font.size = Pt(8.5)
r_cap2.italic = True
r_cap2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

# ==========================================
# 6. Deskripsi Produk
# ==========================================
# Original: P4 = Heading 2 "Deskripsi Produk", P5 = Placeholder
p_desc_h = doc.paragraphs[4]
p_desc_h.paragraph_format.space_before = Pt(10)
p_desc_h.paragraph_format.space_after = Pt(4)

# Replace P5 with Paragraph 1 of Deskripsi Produk
p5 = doc.paragraphs[5]
p5.text = ""
p5.paragraph_format.space_before = Pt(0)
p5.paragraph_format.space_after = Pt(4)
p5.paragraph_format.line_spacing = 1.15
p5.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
r5_1 = p5.add_run("CareerStudent AI adalah platform Adaptive Career Co-Pilot berbasis web (Progressive Web Apps/PWA) yang dirancang khusus untuk mentransformasi kesiapan kerja mahasiswa tingkat akhir dan fresh graduates lintas disiplin keilmuan (Teknologi Informasi, Bisnis Digital, Desain Komunikasi Visual, dan Bahasa Inggris Bisnis). Inovasi ini hadir merespons fenomena maraknya vibe-applying—kebiasaan mahasiswa yang mengandalkan otomasi kecerdasan buatan (generative AI) secara pasif tanpa memahami narasi portofolio mereka sendiri, yang berujung pada tingginya tingkat kegagalan saat melalui tahapan wawancara mendalam.")
r5_1.font.name = "Calibri"
r5_1.font.size = Pt(10)

# Insert Paragraph 2 of Deskripsi Produk before P6 ("Kebermanfaatan Produk")
p5_2 = doc.paragraphs[6].insert_paragraph_before()
p5_2.paragraph_format.space_before = Pt(0)
p5_2.paragraph_format.space_after = Pt(6)
p5_2.paragraph_format.line_spacing = 1.15
p5_2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
r5_2 = p5_2.add_run("Berbeda dengan generator CV konvensional yang mengotomatisasi konten secara instan tanpa proses kognitif, CareerStudent AI mengusung paradigma pedagogis Human-in-the-Loop (HITL) dan Socratic Dialogue. Platform ini memiliki tiga modul inti: (1) ATS-Compliant Fast CV Builder dengan validasi standar industri global, (2) Skill-Gap Matrix berbasis taksonomi World Economic Forum (WEF) 2023 yang memetakan disparitas kompetensi secara riil, serta (3) Interactive Mock Interview dengan kerangka kerja STAR (Situation, Task, Action, Result). Melalui mekanisme Explicit Approval Gate dan Mandatory Self-Reflection Gate, AI bertindak sebagai fasilitator kritis yang menantang nalar mahasiswa, melarang injeksi teks otomatis tanpa persetujuan eksplisit, serta mengunci laporan evaluasi sebelum pengguna merefleksikan pembelajarannya secara mandiri.")
r5_2.font.name = "Calibri"
r5_2.font.size = Pt(10)

# ==========================================
# 7. Kebermanfaatan Produk
# ==========================================
# Now find "Kebermanfaatan Produk" heading
idx_keb = None
for i, p in enumerate(doc.paragraphs):
    if "Kebermanfaatan Produk" in p.text:
        idx_keb = i
        break

p_keb_h = doc.paragraphs[idx_keb]
p_keb_h.paragraph_format.space_before = Pt(10)
p_keb_h.paragraph_format.space_after = Pt(4)

# Body 1 of Kebermanfaatan is idx_keb + 1
p_keb1 = doc.paragraphs[idx_keb + 1]
p_keb1.text = ""
p_keb1.paragraph_format.space_before = Pt(0)
p_keb1.paragraph_format.space_after = Pt(4)
p_keb1.paragraph_format.line_spacing = 1.15
p_keb1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
r_keb1 = p_keb1.add_run("CareerStudent AI memberikan kebermanfaatan nyata bagi mahasiswa multidisiplin dan pengelola Career Development Center (CDC) perguruan tinggi. Platform ini secara efektif menyelesaikan akar masalah vibe-applying dengan mengembalikan kedaulatan berpikir dan daya nalar kritis mahasiswa terhadap portofolionya. Dari aspek efisiensi operasional, CareerStudent AI memangkas waktu penyusunan CV siap kerja berstandar ATS dari rata-rata 3 jam menjadi kurang dari 5 menit tanpa mengorbankan kedalaman konten.")
r_keb1.font.name = "Calibri"
r_keb1.font.size = Pt(10)

# Find "Informasi Singkat" heading
idx_info = None
for i, p in enumerate(doc.paragraphs):
    if "Informasi Singkat" in p.text:
        idx_info = i
        break

p_keb2 = doc.paragraphs[idx_info].insert_paragraph_before()
p_keb2.paragraph_format.space_before = Pt(0)
p_keb2.paragraph_format.space_after = Pt(6)
p_keb2.paragraph_format.line_spacing = 1.15
p_keb2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
r_keb2 = p_keb2.add_run("Dampak terukur telah divalidasi melalui 42 skenario pengujian fungsional black-box dengan tingkat keberhasilan 100% dan lolos stress-testing beban konkurensi tinggi. Pengujian empiris pada kelompok mahasiswa uji coba menunjukkan peningkatan skor kelayakan CV rata-rata sebesar +38% serta retensi pemahaman narasi portofolio saat wawancara mencapai 92%. Inovasi ini membuka peluang karir yang lebih inklusif dan berdaya saing tinggi bagi lulusan perguruan tinggi di era kecerdasan buatan.")
r_keb2.font.name = "Calibri"
r_keb2.font.size = Pt(10)

# ==========================================
# 8. Informasi Singkat
# ==========================================
# Find idx_info again
for i, p in enumerate(doc.paragraphs):
    if "Informasi Singkat" in p.text:
        p_info_h = p
        break

p_info_h.paragraph_format.space_before = Pt(10)
p_info_h.paragraph_format.space_after = Pt(4)

# ==========================================
# 9. TABLE 2: Foto Tim
# ==========================================
t2 = doc.tables[2]
c2 = t2.rows[0].cells[0]
p2_cell = c2.paragraphs[0]
p2_cell.text = ""
p2_cell.paragraph_format.space_before = Pt(8)
p2_cell.paragraph_format.space_after = Pt(4)
p2_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER

r2_c_t = p2_cell.add_run("TIM AQUAMAN TEKNOKRAT\n")
r2_c_t.font.name = "Calibri"
r2_c_t.font.size = Pt(11)
r2_c_t.bold = True
r2_c_t.font.color.rgb = RGBColor(0x17, 0x32, 0x4D)

r2_c_sub = p2_cell.add_run("Universitas Teknokrat Indonesia\n\n")
r2_c_sub.font.name = "Calibri"
r2_c_sub.font.size = Pt(9)
r2_c_sub.font.color.rgb = RGBColor(0x00, 0xA6, 0xA6)
r2_c_sub.bold = True

r2_m1 = p2_cell.add_run("1. Riki Ardiansyah (Ketua Tim / Lead Developer)\n")
r2_m1.font.name = "Calibri"
r2_m1.font.size = Pt(8.5)
r2_m1.bold = True

r2_m2 = p2_cell.add_run("2. Jovan Akbar Setiodono (Anggota / UI-UX Designer)\n")
r2_m2.font.name = "Calibri"
r2_m2.font.size = Pt(8.5)
r2_m2.bold = True

r2_m3 = p2_cell.add_run("3. Fiqi Rahman Al Afkhoni (Anggota / AI & QA Engineer)\n\n")
r2_m3.font.name = "Calibri"
r2_m3.font.size = Pt(8.5)
r2_m3.bold = True

r2_note = p2_cell.add_run("[Area Pas Foto Formal Anggota Tim]")
r2_note.font.name = "Calibri"
r2_note.font.size = Pt(8)
r2_note.italic = True
r2_note.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

# ==========================================
# 10. TABLE 3: Metadata Tim
# ==========================================
t3 = doc.tables[3]

# Row 0: Nama Tim
c_t3_01 = t3.rows[0].cells[1]
c_t3_01.paragraphs[0].text = ""
r_t3_0 = c_t3_01.paragraphs[0].add_run("TIM AQUAMAN TEKNOKRAT")
r_t3_0.font.name = "Calibri"
r_t3_0.font.size = Pt(9)
r_t3_0.bold = True
r_t3_0.font.color.rgb = RGBColor(0x17, 0x32, 0x4D)

# Row 1: Anggota Tim
c_t3_11 = t3.rows[1].cells[1]
c_t3_11.paragraphs[0].text = ""
p_m = c_t3_11.paragraphs[0]
p_m.paragraph_format.space_before = Pt(2)
p_m.paragraph_format.space_after = Pt(2)
p_m.paragraph_format.line_spacing = 1.15

r_m1 = p_m.add_run("1. Riki Ardiansyah (Ketua - 24313003 - S1 Teknologi Informasi)\n")
r_m1.font.name = "Calibri"
r_m1.font.size = Pt(9)
r_m1.bold = True

r_m2 = p_m.add_run("2. Jovan Akbar Setiodono (Anggota - 24313024 - S1 Teknologi Informasi)\n")
r_m2.font.name = "Calibri"
r_m2.font.size = Pt(9)

r_m3 = p_m.add_run("3. Fiqi Rahman Al Afkhoni (Anggota - 24313023 - S1 Teknologi Informasi)")
r_m3.font.name = "Calibri"
r_m3.font.size = Pt(9)

# Row 2: Kategori / Bidang
c_t3_21 = t3.rows[2].cells[1]
c_t3_21.paragraphs[0].text = ""
r_t3_2 = c_t3_21.paragraphs[0].add_run("Artificial Intelligence & Education Technology (EdTech)")
r_t3_2.font.name = "Calibri"
r_t3_2.font.size = Pt(9)

# Row 3: Asal Sekolah / Universitas / Umum
c_t3_31 = t3.rows[3].cells[1]
c_t3_31.paragraphs[0].text = ""
r_t3_3 = c_t3_31.paragraphs[0].add_run("Universitas Teknokrat Indonesia")
r_t3_3.font.name = "Calibri"
r_t3_3.font.size = Pt(9)
r_t3_3.bold = True

# Save to both paths
doc.save(OUTPUT_SCRATCH)
doc.save(OUTPUT_BRAIN)
print(f"Successfully generated:\n- {OUTPUT_SCRATCH}\n- {OUTPUT_BRAIN}")

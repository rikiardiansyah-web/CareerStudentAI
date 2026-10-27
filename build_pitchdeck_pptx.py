import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# -------------------------------------------------------------
# CONSTANTS & PALETTE
# -------------------------------------------------------------
C_NAVY_DARK   = RGBColor(15, 23, 42)     # #0F172A Slate 900
C_NAVY_MED    = RGBColor(23, 50, 77)     # #17324D Navy Accent
C_TEAL        = RGBColor(0, 166, 166)    # #00A6A6 Brand Teal
C_TEAL_LIGHT  = RGBColor(234, 244, 244)  # #EAF4F4
C_CYAN        = RGBColor(14, 165, 233)   # #0EA5E9
C_GOLD        = RGBColor(245, 158, 11)   # #F59E0B Golden Amber
C_BG_LIGHT    = RGBColor(248, 250, 252)  # #F8FAFC
C_CARD_BG     = RGBColor(255, 255, 255)  # #FFFFFF
C_CARD_BORDER = RGBColor(226, 232, 240)  # #E2E8F0
C_TEXT_DARK   = RGBColor(15, 23, 42)     # #0F172A
C_TEXT_MUTED  = RGBColor(100, 116, 139)  # #64748B
C_TEXT_WHITE  = RGBColor(255, 255, 255)  # #FFFFFF

LOGO_PATH = r"C:\Users\Admin\.gemini\antigravity\scratch\careerstudent-ai\public\icons\logo-cs.jpg"
SS1_PATH  = r"C:\Users\Admin\.gemini\antigravity\brain\8c1f2961-a2f7-42ef-86e0-6c72fe3fc1a0\.user_uploaded\media_1792359115071_ae73705e.jpg"
SS2_PATH  = r"C:\Users\Admin\.gemini\antigravity\brain\8c1f2961-a2f7-42ef-86e0-6c72fe3fc1a0\.user_uploaded\media_1792361157177_53c3abec.jpg"

OUTPUT_SCRATCH = r"C:\Users\Admin\.gemini\antigravity\scratch\careerstudent-ai\PitchDeck_CareerStudentAI.pptx"
OUTPUT_BRAIN   = r"C:\Users\Admin\.gemini\antigravity\brain\8c1f2961-a2f7-42ef-86e0-6c72fe3fc1a0\PitchDeck_CareerStudentAI.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# -------------------------------------------------------------
# HELPER FUNCTIONS
# -------------------------------------------------------------
def set_shape_fill(shape, fill_color, border_color=None, border_width=1):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(border_width)
    else:
        shape.line.fill.background()

def add_header(slide, title, category="TIM AQUAMAN TEKNOKRAT  |  GDI FEST 2026", subtitle=None):
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.1))
    set_shape_fill(top_bar, C_TEAL)

    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.35), Inches(4.5), Inches(0.32))
    set_shape_fill(badge, C_TEAL_LIGHT, C_TEAL, border_width=1)
    tf_b = badge.text_frame
    tf_b.word_wrap = True
    p_b = tf_b.paragraphs[0]
    p_b.text = category
    p_b.font.size = Pt(9.5)
    p_b.font.bold = True
    p_b.font.color.rgb = C_NAVY_MED
    p_b.alignment = PP_ALIGN.CENTER

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.7))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = title
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.color.rgb = C_NAVY_DARK

    if subtitle:
        p_sub = tf_t.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = C_TEXT_MUTED

def add_footer(slide, is_dark=False):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.04))
    set_shape_fill(line, C_TEAL if is_dark else C_CARD_BORDER)

    box_l = slide.shapes.add_textbox(Inches(0.8), Inches(7.12), Inches(7.0), Inches(0.3))
    tf_l = box_l.text_frame
    tf_l.margin_left = tf_l.margin_right = tf_l.margin_top = tf_l.margin_bottom = 0
    p_l = tf_l.paragraphs[0]
    p_l.text = "CAREERSTUDENT AI  •  ADAPTIVE CAREER CO-PILOT"
    p_l.font.size = Pt(9)
    p_l.font.bold = True
    p_l.font.color.rgb = C_TEAL if is_dark else C_TEXT_MUTED

    box_r = slide.shapes.add_textbox(Inches(7.0), Inches(7.12), Inches(5.533), Inches(0.3))
    tf_r = box_r.text_frame
    tf_r.margin_left = tf_r.margin_right = tf_r.margin_top = tf_r.margin_bottom = 0
    p_r = tf_r.paragraphs[0]
    p_r.alignment = PP_ALIGN.RIGHT
    run1 = p_r.add_run()
    run1.text = "UNIVERSITAS TEKNOKRAT INDONESIA "
    run1.font.bold = True
    run1.font.size = Pt(9)
    run1.font.color.rgb = C_GOLD if is_dark else C_NAVY_MED
    run2 = p_r.add_run()
    run2.text = "• KAMPUSNYA SANG JUARA"
    run2.font.bold = True
    run2.font.size = Pt(9)
    run2.font.color.rgb = C_TEAL if is_dark else C_TEAL

def create_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_CARD_BORDER, border_width=1):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    set_shape_fill(card, bg_color, border_color, border_width)
    return card

# =============================================================
# SLIDE 1: COVER / TITLE SLIDE (Dark Navy Theme)
# =============================================================
slide1 = prs.slides.add_slide(blank_layout)
bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
set_shape_fill(bg1, C_NAVY_DARK)

bar1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
set_shape_fill(bar1, C_TEAL)

badge1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(4.5), Inches(0.42))
set_shape_fill(badge1, RGBColor(30, 41, 59), C_TEAL, border_width=1.5)
tf_b1 = badge1.text_frame
p_b1 = tf_b1.paragraphs[0]
p_b1.text = "GDI FEST 2026  |  LOMBA INOVASI DIGITAL"
p_b1.font.size = Pt(11)
p_b1.font.bold = True
p_b1.font.color.rgb = C_TEAL
p_b1.alignment = PP_ALIGN.CENTER

title_box1 = slide1.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(8.5), Inches(2.2))
tf1 = title_box1.text_frame
tf1.word_wrap = True
tf1.margin_left = tf1.margin_right = tf1.margin_top = tf1.margin_bottom = 0

p1_sub = tf1.paragraphs[0]
p1_sub.text = "PITCH DECK INOVASI PENDIDIKAN & TEKNOLOGI"
p1_sub.font.size = Pt(12)
p1_sub.font.bold = True
p1_sub.font.color.rgb = C_GOLD

p1_main = tf1.add_paragraph()
p1_main.text = "CAREERSTUDENT AI"
p1_main.font.size = Pt(44)
p1_main.font.bold = True
p1_main.font.color.rgb = C_TEXT_WHITE

p1_tag = tf1.add_paragraph()
p1_tag.text = "Adaptive Career Co-Pilot Berbasis Human-in-the-Loop & Socratic Dialogue untuk Menjembatani Kesiapan Kerja Mahasiswa Multidisiplin di Era AI Generatif"
p1_tag.font.size = Pt(14)
p1_tag.font.color.rgb = RGBColor(203, 213, 225)

if os.path.exists(LOGO_PATH):
    logo_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.8), Inches(1.4), Inches(2.7), Inches(2.7))
    set_shape_fill(logo_card, RGBColor(255, 255, 255), C_TEAL, border_width=2)
    slide1.shapes.add_picture(LOGO_PATH, Inches(10.05), Inches(1.65), width=Inches(2.2))

team_bar = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.3), Inches(11.733), Inches(2.2))
set_shape_fill(team_bar, RGBColor(23, 37, 65), C_TEAL, border_width=1)
tf_tb = team_bar.text_frame
tf_tb.word_wrap = True

p_t1 = tf_tb.paragraphs[0]
p_t1.text = "DIKEMBANGKAN OLEH :"
p_t1.font.size = Pt(10)
p_t1.font.bold = True
p_t1.font.color.rgb = C_TEAL

p_t2 = tf_tb.add_paragraph()
p_t2.text = "TIM AQUAMAN TEKNOKRAT"
p_t2.font.size = Pt(20)
p_t2.font.bold = True
p_t2.font.color.rgb = C_TEXT_WHITE

p_t3 = tf_tb.add_paragraph()
p_t3.text = "• Riki Ardiansyah (Ketua / Lead Developer - 24313003)    • Jovan Akbar Setiodono (UI/UX - 24313024)    • Fiqi Rahman Al Afkhoni (AI/QA - 24313023)"
p_t3.font.size = Pt(11)
p_t3.font.color.rgb = RGBColor(226, 232, 240)

p_t4 = tf_tb.add_paragraph()
p_t4.text = "UNIVERSITAS TEKNOKRAT INDONESIA — KAMPUSNYA SANG JUARA"
p_t4.font.size = Pt(12)
p_t4.font.bold = True
p_t4.font.color.rgb = C_GOLD

add_footer(slide1, is_dark=True)


# =============================================================
# SLIDE 2: AQUAMAN TEAM PROFILE
# =============================================================
slide2 = prs.slides.add_slide(blank_layout)
add_header(slide2, "Profil Tim Pengembang: AQUAMAN Team", subtitle="Kolaborasi Mahasiswa Multidisiplin Teknologi Informasi — Universitas Teknokrat Indonesia")

members = [
    {
        "role_badge": "KETUA TIM / LEAD ARCHITECT",
        "name": "Riki Ardiansyah",
        "nim": "NIM: 24313003",
        "prodi": "S1 Teknologi Informasi",
        "focus": [
            "Full-Stack Web Architecture (Next.js 16 PWA)",
            "High-Performance Microservices (Golang gRPC)",
            "Socratic Dialogue & LLM Prompt Engineering",
            "Security Hardening (Scrypt & OWASP Compliance)"
        ],
        "accent": C_TEAL
    },
    {
        "role_badge": "ANGGOTA / UI-UX SPECIALIST",
        "name": "Jovan Akbar Setiodono",
        "nim": "NIM: 24313024",
        "prodi": "S1 Teknologi Informasi",
        "focus": [
            "Human-Centered UI/UX Design System",
            "Socratic Interaction Flow & Usability Testing",
            "Multidisciplinary CV Presets Adaptation",
            "Responsive Web & Mobile Layout Integration"
        ],
        "accent": C_CYAN
    },
    {
        "role_badge": "ANGGOTA / AI & QA ENGINEER",
        "name": "Fiqi Rahman Al Afkhoni",
        "nim": "NIM: 24313023",
        "prodi": "S1 Teknologi Informasi",
        "focus": [
            "STAR Interview Simulation & Web Speech Engine",
            "Skill-Gap Matrix (WEF 2023 Taxonomy)",
            "42 Black-box Functional Test Automation",
            "User Readiness & Cognitive Feedback Analytics"
        ],
        "accent": C_GOLD
    }
]

for idx, m in enumerate(members):
    left = 0.8 + idx * 4.0
    top = 1.6
    card = create_card(slide2, left, top, 3.733, 5.15, bg_color=C_CARD_BG, border_color=C_CARD_BORDER, border_width=1.5)
    
    badge_sh = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left + 0.2), Inches(top + 0.2), Inches(3.333), Inches(0.35))
    set_shape_fill(badge_sh, m["accent"])
    p_mb = badge_sh.text_frame.paragraphs[0]
    p_mb.text = m["role_badge"]
    p_mb.font.size = Pt(8.5)
    p_mb.font.bold = True
    p_mb.font.color.rgb = C_TEXT_WHITE
    p_mb.alignment = PP_ALIGN.CENTER
    
    m_box = slide2.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.65), Inches(3.333), Inches(4.3))
    tf_m = m_box.text_frame
    tf_m.word_wrap = True
    tf_m.margin_left = tf_m.margin_right = tf_m.margin_top = tf_m.margin_bottom = 0
    
    p_name = tf_m.paragraphs[0]
    p_name.text = m["name"]
    p_name.font.size = Pt(17)
    p_name.font.bold = True
    p_name.font.color.rgb = C_NAVY_DARK
    
    p_nim = tf_m.add_paragraph()
    p_nim.text = f"{m['nim']} • {m['prodi']}"
    p_nim.font.size = Pt(10)
    p_nim.font.bold = True
    p_nim.font.color.rgb = C_TEXT_MUTED
    
    p_div = tf_m.add_paragraph()
    p_div.text = "————————————————————"
    p_div.font.size = Pt(8)
    p_div.font.color.rgb = C_CARD_BORDER
    
    p_foc_lbl = tf_m.add_paragraph()
    p_foc_lbl.text = "Fokus & Tanggung Jawab Teknis:"
    p_foc_lbl.font.size = Pt(10.5)
    p_foc_lbl.font.bold = True
    p_foc_lbl.font.color.rgb = C_NAVY_MED
    
    for f in m["focus"]:
        pf = tf_m.add_paragraph()
        pf.text = f"✓ {f}"
        pf.font.size = Pt(9.5)
        pf.font.color.rgb = C_TEXT_DARK

add_footer(slide2)


# =============================================================
# SLIDE 3: LATAR BELAKANG (PROBLEM BACKGROUND)
# =============================================================
slide3 = prs.slides.add_slide(blank_layout)
add_header(slide3, "Latar Belakang: Krisis Kesiapan Kerja & Paradoks AI Generatif", 
           subtitle="Fenomena 'Vibe-Applying' Mengancam Kedaulatan Kognitif Mahasiswa Pencari Kerja")

bg_ctx = create_card(slide3, 0.8, 1.55, 11.733, 1.0, bg_color=C_TEAL_LIGHT, border_color=C_TEAL, border_width=1)
tf_ctx = bg_ctx.text_frame
tf_ctx.word_wrap = True
p_c1 = tf_ctx.paragraphs[0]
p_c1.text = "Konteks Kritis Era Disrupsi AI:"
p_c1.font.size = Pt(10.5)
p_c1.font.bold = True
p_c1.font.color.rgb = C_NAVY_DARK

p_c2 = tf_ctx.add_paragraph()
p_c2.text = "Generative AI mempermudah pembuatan dokumen karir, namun justru memicu fenomena 'vibe-applying'—mahasiswa mengandalkan otomatisasi copy-paste tanpa memahami substansi portofolionya sendiri. Dampaknya, kandidat gugur massal saat melalui uji wawancara mendalam."
p_c2.font.size = Pt(10.5)
p_c2.font.color.rgb = C_NAVY_MED

lb_points = [
    {
        "num": "01",
        "title": "Fenomena Vibe-Applying & Copy-Paste Pasif",
        "desc": "Mahasiswa mempercayakan seluruh penulisan CV ke LLM publik tanpa proses verifikasi dan refleksi, menciptakan portofolio generik yang rawan halusinasi.",
        "color": C_TEAL
    },
    {
        "num": "02",
        "title": "Paradoks Wawancara Mendalam (Interview Freeze)",
        "desc": "CV tampak impresif di atas kertas, namun pelamar gagal total ketika recruiter menanyakan studi kasus spesifik dan metodologi pemecahan masalah.",
        "color": C_CYAN
    },
    {
        "num": "03",
        "title": "Kesenjangan Kompetensi Multidisiplin",
        "desc": "Mahasiswa dari 4 klaster keilmuan (TI, Bisnis, DKV, Bahasa) kesulitan memetakan kompetensi riil mereka dengan standar ATS industri global.",
        "color": C_GOLD
    },
    {
        "num": "04",
        "title": "Degradasi Daya Nalar Kritis (Cognitive Atrophy)",
        "desc": "Ketiadaan dialog sokratik menghilangkan ruang belajar reflektif, memperlebar jurang antara kualifikasi ijazah dan kesiapan mental dunia kerja riil.",
        "color": RGBColor(239, 68, 68)
    }
]

for idx, pt in enumerate(lb_points):
    col = idx % 2
    row = idx // 2
    left = 0.8 + col * 6.0
    top = 2.75 + row * 2.05
    
    card = create_card(slide3, left, top, 5.733, 1.85, bg_color=C_CARD_BG, border_color=C_CARD_BORDER)
    
    pill = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left + 0.2), Inches(top + 0.2), Inches(0.65), Inches(0.4))
    set_shape_fill(pill, pt["color"])
    p_num = pill.text_frame.paragraphs[0]
    p_num.text = pt["num"]
    p_num.font.size = Pt(11)
    p_num.font.bold = True
    p_num.font.color.rgb = C_TEXT_WHITE
    p_num.alignment = PP_ALIGN.CENTER
    
    tb = slide3.shapes.add_textbox(Inches(left + 0.95), Inches(top + 0.15), Inches(4.6), Inches(1.55))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    
    p_t = tf.paragraphs[0]
    p_t.text = pt["title"]
    p_t.font.size = Pt(12)
    p_t.font.bold = True
    p_t.font.color.rgb = C_NAVY_DARK
    
    p_d = tf.add_paragraph()
    p_d.text = pt["desc"]
    p_d.font.size = Pt(10)
    p_d.font.color.rgb = C_TEXT_MUTED

add_footer(slide3)


# =============================================================
# SLIDE 4: RUANG LINGKUP PROYEK SISTEM (SYSTEM SCOPE)
# =============================================================
slide4 = prs.slides.add_slide(blank_layout)
add_header(slide4, "Ruang Lingkup Proyek Sistem", subtitle="5 Pilar Inovasi Kognitif & Fungsional CareerStudent AI")

scopes = [
    {
        "num": "01",
        "title": "ATS-Compliant Fast CV Builder (4 Preset Multidisiplin)",
        "desc": "Generator resume standar ATS internasional yang dikhususkan untuk 4 program studi: Teknologi Informasi, Bisnis Digital, Desain Komunikasi Visual (DKV), dan Bahasa Inggris Bisnis.",
        "tag": "Modul Pembuatan Karir"
    },
    {
        "num": "02",
        "title": "Explicit Approval Gate (Human-in-the-Loop Safeguard)",
        "desc": "Mekanisme penjaga pedagogis yang melarang AI menulis atau menyunting CV secara otonom tanpa adanya peninjauan, persetujuan eksplisit, dan validasi sadar dari mahasiswa.",
        "tag": "Filosofi Kedaulatan Kognitif"
    },
    {
        "num": "03",
        "title": "Socratic Dialogue Coaching Engine",
        "desc": "Mesin dialog berbasis pertanyaan sokratik terstruktur yang menantang mahasiswa untuk menggali metodologi, metrik kuantitatif, dan kedalaman peran dari setiap pengalaman mereka.",
        "tag": "Penalaran Kritis"
    },
    {
        "num": "04",
        "title": "Skill-Gap Matrix Berbasis Taksonomi WEF 2023",
        "desc": "Modul diagnostik otomatis yang memetakan kualifikasi pengguna terhadap standar 10 keterampilan masa depan World Economic Forum (WEF) 2023 untuk mendeteksi kesenjangan skill.",
        "tag": "Analitik Kompetensi"
    },
    {
        "num": "05",
        "title": "Interactive Mock Interview STAR & Mandatory Reflection Gate",
        "desc": "Simulasi wawancara audio-visual interaktif dengan kerangka kerja STAR (Situation, Task, Action, Result). Hasil evaluasi dikunci hingga kandidat menyelesaikan refleksi mandiri.",
        "tag": "Kesiapan Mental & Wawancara"
    }
]

for idx, sc in enumerate(scopes):
    top = 1.6 + idx * 1.05
    card = create_card(slide4, 0.8, top, 11.733, 0.95, bg_color=C_CARD_BG, border_color=C_CARD_BORDER)
    
    pill = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(top + 0.2), Inches(0.7), Inches(0.55))
    set_shape_fill(pill, C_TEAL)
    p_num = pill.text_frame.paragraphs[0]
    p_num.text = sc["num"]
    p_num.font.size = Pt(13)
    p_num.font.bold = True
    p_num.font.color.rgb = C_TEXT_WHITE
    p_num.alignment = PP_ALIGN.CENTER
    
    tb = slide4.shapes.add_textbox(Inches(1.85), Inches(top + 0.12), Inches(8.3), Inches(0.75))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    
    p_t = tf.paragraphs[0]
    p_t.text = sc["title"]
    p_t.font.size = Pt(12)
    p_t.font.bold = True
    p_t.font.color.rgb = C_NAVY_DARK
    
    p_d = tf.add_paragraph()
    p_d.text = sc["desc"]
    p_d.font.size = Pt(9.5)
    p_d.font.color.rgb = C_TEXT_MUTED
    
    tag_sh = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.3), Inches(top + 0.3), Inches(2.0), Inches(0.35))
    set_shape_fill(tag_sh, C_TEAL_LIGHT, C_TEAL, border_width=1)
    p_tg = tag_sh.text_frame.paragraphs[0]
    p_tg.text = sc["tag"]
    p_tg.font.size = Pt(8.5)
    p_tg.font.bold = True
    p_tg.font.color.rgb = C_NAVY_MED
    p_tg.alignment = PP_ALIGN.CENTER

add_footer(slide4)


# =============================================================
# SLIDE 5: SASARAN PROYEK (PROJECT OBJECTIVES)
# =============================================================
slide5 = prs.slides.add_slide(blank_layout)
add_header(slide5, "Sasaran & Target Capaian Proyek", subtitle="Mewujudkan Lulusan yang Mandiri Berpikir, Kompeten, dan Teruji di Dunia Industri")

targets = [
    {
        "num": "01",
        "title": "Mengembalikan Daya Nalar Kritis Mahasiswa",
        "desc": "Menghilangkan ketergantungan pasif pada AI generatif dan memastikan kandidat 100% memahami narasi portofolio yang diajukannya ke perusahaan.",
        "metric": "Target: 90%+ Pemahaman Narasi"
    },
    {
        "num": "02",
        "title": "Meningkatkan Efisiensi Pembuatan Resume ATS",
        "desc": "Memangkas waktu penyusunan CV siap kirim dari rata-rata 3 jam menjadi kurang dari 5 menit tanpa mengorbankan kualitas dan orisinalitas.",
        "metric": "Efisiensi: < 5 Menit / CV"
    },
    {
        "num": "03",
        "title": "Meningkatkan Lolos Screening ATS & Skor Kualifikasi",
        "desc": "Menjamin 100% format dokumen lolos parsing Applicant Tracking System internasional dengan peningkatan relevansi kata kunci industri.",
        "metric": "Skor ATS: Peningkatan ≥ 35%"
    },
    {
        "num": "04",
        "title": "Akselerasi Kesiapan Wawancara Kerja Terstandarisasi",
        "desc": "Membiasakan mahasiswa menyusun jawaban terstruktur dengan metodologi STAR (Situation, Task, Action, Result) sebelum wawancara riil.",
        "metric": "Metode: 100% STAR Framework"
    },
    {
        "num": "05",
        "title": "Menyediakan Solusi Karir Inklusif & Zero-Cost",
        "desc": "Memberikan bimbingan karir setara konsultan profesional kepada seluruh mahasiswa tanpa beban biaya langganan yang mahal.",
        "metric": "Akses: PWA Gratis & Terbuka"
    }
]

for idx, tr in enumerate(targets):
    top = 1.6 + idx * 1.05
    card = create_card(slide5, 0.8, top, 11.733, 0.95, bg_color=C_CARD_BG, border_color=C_CARD_BORDER)
    
    pill = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(top + 0.2), Inches(0.7), Inches(0.55))
    set_shape_fill(pill, C_CYAN)
    p_num = pill.text_frame.paragraphs[0]
    p_num.text = tr["num"]
    p_num.font.size = Pt(13)
    p_num.font.bold = True
    p_num.font.color.rgb = C_TEXT_WHITE
    p_num.alignment = PP_ALIGN.CENTER
    
    tb = slide5.shapes.add_textbox(Inches(1.85), Inches(top + 0.12), Inches(8.0), Inches(0.75))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    
    p_t = tf.paragraphs[0]
    p_t.text = tr["title"]
    p_t.font.size = Pt(12)
    p_t.font.bold = True
    p_t.font.color.rgb = C_NAVY_DARK
    
    p_d = tf.add_paragraph()
    p_d.text = tr["desc"]
    p_d.font.size = Pt(9.5)
    p_d.font.color.rgb = C_TEXT_MUTED
    
    m_badge = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.0), Inches(top + 0.3), Inches(2.3), Inches(0.35))
    set_shape_fill(m_badge, RGBColor(254, 243, 199), C_GOLD, border_width=1)
    p_mb = m_badge.text_frame.paragraphs[0]
    p_mb.text = tr["metric"]
    p_mb.font.size = Pt(8.5)
    p_mb.font.bold = True
    p_mb.font.color.rgb = RGBColor(180, 83, 9)
    p_mb.alignment = PP_ALIGN.CENTER

add_footer(slide5)


# =============================================================
# SLIDE 6: PERMASALAHAN-PERMASALAHAN (DEEP DIVE PAIN POINTS)
# =============================================================
slide6 = prs.slides.add_slide(blank_layout)
add_header(slide6, "Permasalahan Kritis Yang Dipecahkan", subtitle="Analisis Akar Masalah yang Dihadapi Mahasiswa dan Dunia Industri Saat Ini")

problems = [
    {
        "num": "01",
        "title": "Otomasi Tanpa Pemahaman (Surface-Level AI)",
        "desc": "Mahasiswa menyalin teks LLM mentah-mentah ke CV tanpa memahami logika di balik klaim kompetensi tersebut, memicu kegagalan saat wawancara."
    },
    {
        "num": "02",
        "title": "Tingginya Angka Penolakan Parser ATS",
        "desc": "Mayoritas template CV mahasiswa menggunakan grafis, tabel ganda, dan ikon yang tidak terbaca oleh sistem ATS perusahaan besar."
    },
    {
        "num": "03",
        "title": "Sindrom Pembekuan Wawancara (Interview Freeze)",
        "desc": "Kandidat gugup dan gagal menjelaskan tantangan nyata di lapangan karena portofolio mereka dibuat secara instan oleh AI tanpa proses perenungan."
    },
    {
        "num": "04",
        "title": "Ketiadaan Standar Karir Lintas Rumpun Ilmu",
        "desc": "Lulusan Bisnis, DKV, dan Bahasa kerap dipaksa mengikuti format teknis TI yang kaku tanpa taksonomi kompetensi yang relevan dengan bidangnya."
    },
    {
        "num": "05",
        "title": "Ketidakmampuan Mengukur Kesenjangan Skill",
        "desc": "Mahasiswa tidak mengetahui secara objektif apa saja hard skill dan soft skill yang masih kurang terhadap target lowongan kerja yang diincar."
    },
    {
        "num": "06",
        "title": "Mahalnya Layanan Konsultasi Karir Konvensional",
        "desc": "Jasa review CV dan mock interview profesional mematok tarif ratusan ribu hingga jutaan rupiah, di luar jangkauan kantong mahasiswa."
    }
]

for idx, pr in enumerate(problems):
    col = idx % 2
    row = idx // 2
    left = 0.8 + col * 6.0
    top = 1.6 + row * 1.75
    
    card = create_card(slide6, left, top, 5.733, 1.6, bg_color=C_CARD_BG, border_color=C_CARD_BORDER)
    
    pill = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left + 0.2), Inches(top + 0.2), Inches(0.65), Inches(0.4))
    set_shape_fill(pill, RGBColor(239, 68, 68))
    p_num = pill.text_frame.paragraphs[0]
    p_num.text = pr["num"]
    p_num.font.size = Pt(11)
    p_num.font.bold = True
    p_num.font.color.rgb = C_TEXT_WHITE
    p_num.alignment = PP_ALIGN.CENTER
    
    tb = slide6.shapes.add_textbox(Inches(left + 0.95), Inches(top + 0.15), Inches(4.6), Inches(1.35))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    
    p_t = tf.paragraphs[0]
    p_t.text = pr["title"]
    p_t.font.size = Pt(11.5)
    p_t.font.bold = True
    p_t.font.color.rgb = C_NAVY_DARK
    
    p_d = tf.add_paragraph()
    p_d.text = pr["desc"]
    p_d.font.size = Pt(9.5)
    p_d.font.color.rgb = C_TEXT_MUTED

add_footer(slide6)


# =============================================================
# SLIDE 7: KEBUTUHAN INFORMASI PEMAKAI SISTEM (USER PERSONA)
# =============================================================
slide7 = prs.slides.add_slide(blank_layout)
add_header(slide7, "Kebutuhan Informasi Pemakai Sistem", subtitle="Memetakan Informasi & Fitur Esensial Bagi Dua Pemangku Kepentingan Utama")

card_mhs = create_card(slide7, 0.8, 1.6, 5.733, 5.15, bg_color=C_CARD_BG, border_color=C_TEAL, border_width=1.5)
header_mhs = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.8), Inches(5.333), Inches(0.45))
set_shape_fill(header_mhs, C_TEAL)
p_hm = header_mhs.text_frame.paragraphs[0]
p_hm.text = "1. MAHASISWA MULTIDISIPLIN (END USER)"
p_hm.font.size = Pt(11)
p_hm.font.bold = True
p_hm.font.color.rgb = C_TEXT_WHITE
p_hm.alignment = PP_ALIGN.CENTER

tb_mhs = slide7.shapes.add_textbox(Inches(1.1), Inches(2.35), Inches(5.133), Inches(4.2))
tf_mhs = tb_mhs.text_frame
tf_mhs.word_wrap = True
tf_mhs.margin_left = tf_mhs.margin_right = tf_mhs.margin_top = tf_mhs.margin_bottom = 0

mhs_needs = [
    ("Panduan Formulasi Narasi STAR", "Bantuan sokratik dalam menyusun pengalaman kerja berbasis Situation, Task, Action, dan Result yang terukur."),
    ("Skor Kelayakan ATS & Kesesuaian Keyword", "Evaluasi instan keterbacaan CV oleh mesin parser dan kecocokan kata kunci dengan lowongan target."),
    ("Peta Disparitas Kompetensi (Skill-Gap)", "Visualisasi kesenjangan keterampilan teknis dan non-teknis terhadap standar industri WEF 2023."),
    ("Transkrip & Analitik Performa Wawancara", "Umpan balik komprehensif mengenai artikulasi, kejelasan jawaban, dan kepercayaan diri setelah sesi simulasi."),
    ("Preset Kurikulum 4 Program Studi", "Template format resume yang telah disesuaikan dengan karakteristik TI, Bisnis, DKV, dan Bahasa.")
]

for title, desc in mhs_needs:
    p_t = tf_mhs.add_paragraph() if tf_mhs.paragraphs[0].text else tf_mhs.paragraphs[0]
    p_t.text = f"• {title}"
    p_t.font.size = Pt(11)
    p_t.font.bold = True
    p_t.font.color.rgb = C_NAVY_DARK
    
    p_d = tf_mhs.add_paragraph()
    p_d.text = f"   {desc}"
    p_d.font.size = Pt(9.5)
    p_d.font.color.rgb = C_TEXT_MUTED

card_cdc = create_card(slide7, 6.8, 1.6, 5.733, 5.15, bg_color=C_CARD_BG, border_color=C_CYAN, border_width=1.5)
header_cdc = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(1.8), Inches(5.333), Inches(0.45))
set_shape_fill(header_cdc, C_CYAN)
p_hc = header_cdc.text_frame.paragraphs[0]
p_hc.text = "2. CAREER DEVELOPMENT CENTER & RECRUITER"
p_hc.font.size = Pt(11)
p_hc.font.bold = True
p_hc.font.color.rgb = C_TEXT_WHITE
p_hc.alignment = PP_ALIGN.CENTER

tb_cdc = slide7.shapes.add_textbox(Inches(7.1), Inches(2.35), Inches(5.133), Inches(4.2))
tf_cdc = tb_cdc.text_frame
tf_cdc.word_wrap = True
tf_cdc.margin_left = tf_cdc.margin_right = tf_cdc.margin_top = tf_cdc.margin_bottom = 0

cdc_needs = [
    ("Dashboard Agregat Kesiapan Kerja Lulusan", "Data analitik kesiapan kerja mahasiswa tingkat akhir lintas fakultas untuk evaluasi kurikulum prodi."),
    ("Rekapitulasi Distribusi Skill Standar WEF", "Laporan statistik mengenai penguasaan 10 keterampilan utama masa depan mahasiswa kampus mitra."),
    ("Verifikasi Keaslian & Integritas Portofolio", "Jaminan bahwa portofolio kandidat bukan sekadar halusinasi AI, melainkan hasil pemahaman mandiri kandidat."),
    ("Efisiensi Pra-Screening Kandidat", "Mempermudah recruiter menemukan kandidat dengan kualifikasi yang relevan dan teruji pemahamannya."),
    ("Toleransi Akses Luas via Standar Web PWA", "Dapat diakses instan oleh seluruh sivitas akademika tanpa memerlukan instalasi aplikasi native rumit.")
]

for title, desc in cdc_needs:
    p_t = tf_cdc.add_paragraph() if tf_cdc.paragraphs[0].text else tf_cdc.paragraphs[0]
    p_t.text = f"• {title}"
    p_t.font.size = Pt(11)
    p_t.font.bold = True
    p_t.font.color.rgb = C_NAVY_DARK
    
    p_d = tf_cdc.add_paragraph()
    p_d.text = f"   {desc}"
    p_d.font.size = Pt(9.5)
    p_d.font.color.rgb = C_TEXT_MUTED

add_footer(slide7)


# =============================================================
# SLIDE 8: KEGIATAN PENGEMBANGAN YANG TELAH DILAKUKAN
# =============================================================
slide8 = prs.slides.add_slide(blank_layout)
add_header(slide8, "Kegiatan Pengembangan Yang Telah Dilakukan", subtitle="7 Tahapan Rekayasa Sistem Perangkat Lunak dari Riset Kognitif Hingga Deployment Global")

activities = [
    ("1", "Analisis Pedagogis & Taksonomi Industri", "Membedah fenomena vibe-applying, standar parser ATS global, dan 10 core skills laporan WEF 2023."),
    ("2", "Arsitektur Next.js 16 (App Router) & UI PWA", "Membangun sistem antarmuka berbasis React 19, Tailwind CSS, Lucide icons, dan Progressive Web Apps."),
    ("3", "Rekayasa Socratic Engine & Approval Gate", "Merancang alur prompt AI sokratik dan memprogram mekanisme penjaga 'Explicit Approval Gate' anti-otomatisasi."),
    ("4", "Desain Microservices Engine (Golang gRPC)", "Merancang kontrak service komputasi taksonomi skill & kalkulasi ATS berkecepatan tinggi."),
    ("5", "Penguatan Keamanan Data & OWASP Hardening", "Implementasi Scrypt password hashing, CSPRNG OTP, HTTP security headers, dan IndexedDB local-first storage."),
    ("6", "Pengujian Fungsional Black-Box (42 Cases)", "Eksekusi 42 skenario pengujian komprehensif pada auth, builder, matrix, dan interview (Tingkat kelulusan 100%)."),
    ("7", "Deployment Produksi Serverless di Vercel", "Rilis live di https://career-student-ai.vercel.app dengan 15 rute aplikasi terverifikasi aktif.")
]

for idx, (num, title, desc) in enumerate(activities):
    top = 1.6 + idx * 0.75
    card = create_card(slide8, 0.8, top, 11.733, 0.68, bg_color=C_CARD_BG, border_color=C_CARD_BORDER)
    
    pill = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(top + 0.12), Inches(0.45), Inches(0.44))
    set_shape_fill(pill, C_TEAL)
    p_n = pill.text_frame.paragraphs[0]
    p_n.text = num
    p_n.font.size = Pt(11)
    p_n.font.bold = True
    p_n.font.color.rgb = C_TEXT_WHITE
    p_n.alignment = PP_ALIGN.CENTER
    
    tb = slide8.shapes.add_textbox(Inches(1.6), Inches(top + 0.08), Inches(10.7), Inches(0.55))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    
    p_t = tf.paragraphs[0]
    r_t = p_t.add_run()
    r_t.text = f"{title} — "
    r_t.font.bold = True
    r_t.font.size = Pt(11)
    r_t.font.color.rgb = C_NAVY_DARK
    
    r_d = p_t.add_run()
    r_d.text = desc
    r_d.font.size = Pt(10)
    r_d.font.color.rgb = C_TEXT_MUTED

add_footer(slide8)


# =============================================================
# SLIDE 9: METODE PENGEMBANGAN SISTEM (METHODOLOGY)
# =============================================================
slide9 = prs.slides.add_slide(blank_layout)
add_header(slide9, "Metode Pengembangan: Agile Socratic-HITL Prototype", 
           subtitle="Pendekatan Iteratif Berkelanjutan yang Berpusat pada Validasi Pedagogis dan Pengalaman Pengguna")

method_steps = [
    {
        "step": "Tahap 1",
        "title": "Problem Discovery & Persona Alignment",
        "desc": "Wawancara mendalam dengan mahasiswa tingkat akhir dan reviewer karir untuk memetakan titik kegagalan saat wawancara kerja dan pola ketergantungan AI.",
        "badge": "Identifikasi Masalah"
    },
    {
        "step": "Tahap 2",
        "title": "Socratic Prompt Modeling & Wireframing",
        "desc": "Merumuskan logika dialog sokratik dan wireframe antarmuka dengan fokus pada 'Explicit Approval Gate' agar AI tidak pernah mengambil alih kontrol.",
        "badge": "Desain Kognitif"
    },
    {
        "step": "Tahap 3",
        "title": "Rapid Prototyping (Next.js & PWA)",
        "desc": "Membangun prototipe fungsional berkecepatan tinggi dengan Next.js 16, Tailwind CSS, Local IndexedDB, dan integrasi speech synthesis.",
        "badge": "Implementasi Kode"
    },
    {
        "step": "Tahap 4",
        "title": "Iterative Testing & Security Auditing",
        "desc": "Melakukan 42 skenario uji coba blackbox, evaluasi celah keamanan (OWASP Top 10), dan simulasi stres pada penyimpanan offline.",
        "badge": "Pengujian Ketat"
    },
    {
        "step": "Tahap 5",
        "title": "Deployment & User Feedback Loop",
        "desc": "Peluncuran global pada infrastruktur serverless Vercel serta pengumpulan umpan balik dari pengguna mahasiswa lintas 4 rumpun keilmuan.",
        "badge": "Produksi Global"
    }
]

for idx, ms in enumerate(method_steps):
    left = 0.8 + idx * 2.4
    card = create_card(slide9, left, 1.6, 2.25, 4.4, bg_color=C_CARD_BG, border_color=C_CARD_BORDER)
    
    sh = slide9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(1.6), Inches(2.25), Inches(0.45))
    set_shape_fill(sh, C_NAVY_MED)
    p_sh = sh.text_frame.paragraphs[0]
    p_sh.text = ms["step"]
    p_sh.font.size = Pt(11)
    p_sh.font.bold = True
    p_sh.font.color.rgb = C_TEXT_WHITE
    p_sh.alignment = PP_ALIGN.CENTER
    
    tb = slide9.shapes.add_textbox(Inches(left + 0.15), Inches(2.15), Inches(1.95), Inches(3.7))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    
    p_t = tf.paragraphs[0]
    p_t.text = ms["title"]
    p_t.font.size = Pt(11)
    p_t.font.bold = True
    p_t.font.color.rgb = C_NAVY_DARK
    
    p_d = tf.add_paragraph()
    p_d.text = ms["desc"]
    p_d.font.size = Pt(9)
    p_d.font.color.rgb = C_TEXT_MUTED
    
    badge_bot = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left + 0.15), Inches(5.4), Inches(1.95), Inches(0.35))
    set_shape_fill(badge_bot, C_TEAL_LIGHT, C_TEAL, border_width=1)
    p_bb = badge_bot.text_frame.paragraphs[0]
    p_bb.text = ms["badge"]
    p_bb.font.size = Pt(8)
    p_bb.font.bold = True
    p_bb.font.color.rgb = C_NAVY_MED
    p_bb.alignment = PP_ALIGN.CENTER

sum_banner = create_card(slide9, 0.8, 6.15, 11.733, 0.75, bg_color=C_TEAL_LIGHT, border_color=C_TEAL, border_width=1)
tf_sb = sum_banner.text_frame
tf_sb.word_wrap = True
p_sb = tf_sb.paragraphs[0]
p_sb.text = "Mengapa Metode Ini Dipilih? Memungkinkan siklus penyempurnaan cepat (rapid turnaround) antara respon model sokratik dan kebutuhan psikologis mahasiswa, menjaga relevansi pedagogis tanpa mengorbankan kecepatan rilis."
p_sb.font.size = Pt(9.5)
p_sb.font.color.rgb = C_NAVY_MED

add_footer(slide9)


# =============================================================
# SLIDE 10: KENDALA PROYEK SISTEM & STRATEGI MITIGASI
# =============================================================
slide10 = prs.slides.add_slide(blank_layout)
add_header(slide10, "Kendala Proyek Sistem & Strategi Mitigasi", subtitle="Manajemen Risiko Teknis, Biaya Operasional, dan Integritas Kognitif Pengguna")

constraints = [
    {
        "num": "1",
        "title": "Biaya Token & Ketergantungan API LLM",
        "prob": "Tantangan: Penggunaan API model bahasa komersial skala besar dapat memicu pembengkakan biaya token seiring lonjakan jumlah pengguna.",
        "sol": "Solusi & Mitigasi: Menerapkan arsitektur Local-First Caching berbasis IndexedDB dan kompresi prompt sokratik yang padat. Biaya operasional dasar mendekati $0.",
        "accent": C_GOLD
    },
    {
        "num": "2",
        "title": "Risiko Halusinasi & Fabrikasi Data AI",
        "prob": "Tantangan: Karakteristik generative AI yang rawan menambahkan kualifikasi fiktif atau hiperbolis ke dalam resume pengguna.",
        "sol": "Solusi & Mitigasi: Mekanisme 'Explicit Approval Gate' mengunci konten secara kaku. AI hanya berstatus sebagai penasihat (co-pilot), bukan penulis otonom.",
        "accent": C_TEAL
    },
    {
        "num": "3",
        "title": "Privasi & Keamanan Data Karir Pengguna",
        "prob": "Tantangan: Data portofolio, riwayat pekerjaan, dan kredensial pengguna bersifat sangat sensitif dan rawan terekspos.",
        "sol": "Solusi & Mitigasi: Enkripsi password menggunakan Scrypt dengan 16-byte salt, constant-time compare, OWASP security headers, dan penyimpanan lokal di sisi klien.",
        "accent": C_CYAN
    }
]

for idx, c in enumerate(constraints):
    left = 0.8 + idx * 4.0
    card = create_card(slide10, left, 1.6, 3.733, 5.15, bg_color=C_CARD_BG, border_color=C_CARD_BORDER, border_width=1.5)
    
    tag_sh = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left + 0.2), Inches(1.8), Inches(3.333), Inches(0.4))
    set_shape_fill(tag_sh, c["accent"])
    p_tsh = tag_sh.text_frame.paragraphs[0]
    p_tsh.text = f"KENDALA 0{c['num']}"
    p_tsh.font.size = Pt(10)
    p_tsh.font.bold = True
    p_tsh.font.color.rgb = C_TEXT_WHITE
    p_tsh.alignment = PP_ALIGN.CENTER
    
    tb = slide10.shapes.add_textbox(Inches(left + 0.2), Inches(2.35), Inches(3.333), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    
    p_t = tf.paragraphs[0]
    p_t.text = c["title"]
    p_t.font.size = Pt(13)
    p_t.font.bold = True
    p_t.font.color.rgb = C_NAVY_DARK
    
    p_div = tf.add_paragraph()
    p_div.text = "————————————————————"
    p_div.font.size = Pt(8)
    p_div.font.color.rgb = C_CARD_BORDER
    
    p_pr = tf.add_paragraph()
    p_pr.text = c["prob"]
    p_pr.font.size = Pt(10)
    p_pr.font.color.rgb = RGBColor(185, 28, 28)
    
    p_sp = tf.add_paragraph()
    p_sp.text = ""
    p_sp.font.size = Pt(4)
    
    p_sl = tf.add_paragraph()
    p_sl.text = c["sol"]
    p_sl.font.size = Pt(10)
    p_sl.font.color.rgb = RGBColor(15, 118, 110)

add_footer(slide10)


# =============================================================
# SLIDE 11: ARSITEKTUR TEKNOLOGI PENGEMBANGAN (INCLUDING GOLANG)
# =============================================================
slide11 = prs.slides.add_slide(blank_layout)
add_header(slide11, "Arsitektur Teknologi & Efisiensi Biaya Sistem", 
           subtitle="Struktur Teknologi Modern Terdistribusi: Next.js 16, Microservices Golang gRPC, dan PWA Offline-First")

tech_items = [
    {
        "cat": "FRONTEND & APP ROUTER",
        "tech": "Next.js 16 (React 19) + TypeScript",
        "desc": "Server-side rendering super cepat, routing modular 15 rute, dan type-safety penuh yang menjamin keandalan sistem.",
        "cost": "Open Source ($0)"
    },
    {
        "cat": "HIGH-PERFORMANCE BACKEND",
        "tech": "Golang (gRPC-Go) & Protobuf v3",
        "desc": "Microservices engine performa tinggi untuk komputasi taksonomi skill WEF 2023 dan parsing skor ATS sub-milidetik.",
        "cost": "High Concurrency ($0)"
    },
    {
        "cat": "OFFLINE & LOCAL STORAGE",
        "tech": "Progressive Web Apps + IndexedDB",
        "desc": "Dapat diinstal di smartphone/desktop, berjalan offline-first, dan data CV tersimpan privat di perangkat lokal pengguna.",
        "cost": "Client Storage ($0)"
    },
    {
        "cat": "KEAMANAN & ENKRIPSI",
        "tech": "Scrypt Hashing & OWASP Headers",
        "desc": "Salt 16-byte dengan verifikasi constant-time melawan timing attacks, CSPRNG OTP, dan HTTP security headers ketat.",
        "cost": "In-House Secure ($0)"
    },
    {
        "cat": "CLOUD & SERVERLESS HOSTING",
        "tech": "Vercel Serverless Edge Platform",
        "desc": "Infrastruktur cloud global dengan otomatisasi CI/CD dari repositori GitHub, high availability 99.9% tanpa cold start.",
        "cost": "Free Tier Live ($0)"
    },
    {
        "cat": "ORM & DATA PERSISTENCE",
        "tech": "Prisma ORM & PostgreSQL / SQLite",
        "desc": "Manajemen skema database terstruktur dengan migrasi aman, mendukung transisi fleksibel lokal ke enterprise cloud.",
        "cost": "Scalable DB ($0)"
    }
]

for idx, ti in enumerate(tech_items):
    col = idx % 2
    row = idx // 2
    left = 0.8 + col * 6.0
    top = 1.6 + row * 1.75
    
    card = create_card(slide11, left, top, 5.733, 1.6, bg_color=C_CARD_BG, border_color=C_CARD_BORDER)
    
    badge = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left + 0.2), Inches(top + 0.15), Inches(2.8), Inches(0.35))
    set_shape_fill(badge, C_NAVY_DARK)
    p_b = badge.text_frame.paragraphs[0]
    p_b.text = ti["cat"]
    p_b.font.size = Pt(8.5)
    p_b.font.bold = True
    p_b.font.color.rgb = C_TEXT_WHITE
    p_b.alignment = PP_ALIGN.CENTER
    
    c_badge = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left + 3.8), Inches(top + 0.15), Inches(1.7), Inches(0.35))
    set_shape_fill(c_badge, C_TEAL_LIGHT, C_TEAL, border_width=1)
    p_cb = c_badge.text_frame.paragraphs[0]
    p_cb.text = ti["cost"]
    p_cb.font.size = Pt(8.5)
    p_cb.font.bold = True
    p_cb.font.color.rgb = C_NAVY_MED
    p_cb.alignment = PP_ALIGN.CENTER
    
    tb = slide11.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.58), Inches(5.333), Inches(0.95))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    
    p_t = tf.paragraphs[0]
    p_t.text = ti["tech"]
    p_t.font.size = Pt(11.5)
    p_t.font.bold = True
    p_t.font.color.rgb = C_NAVY_DARK
    
    p_d = tf.add_paragraph()
    p_d.text = ti["desc"]
    p_d.font.size = Pt(9.5)
    p_d.font.color.rgb = C_TEXT_MUTED

add_footer(slide11)


# =============================================================
# SLIDE 12: MANFAAT PENGEMBANGAN SISTEM (PRODUCT BENEFITS)
# =============================================================
slide12 = prs.slides.add_slide(blank_layout)
add_header(slide12, "Manfaat Nyata Pengembangan Sistem", subtitle="Dampak Positif bagi Mahasiswa Pencari Kerja dan Perguruan Tinggi / CDC Kampus")

card_m = create_card(slide12, 0.8, 1.6, 5.733, 5.15, bg_color=C_CARD_BG, border_color=C_TEAL, border_width=1.5)
header_m = slide12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.8), Inches(5.333), Inches(0.45))
set_shape_fill(header_m, C_TEAL)
p_hm2 = header_m.text_frame.paragraphs[0]
p_hm2.text = "1. MANFAAT BAGI MAHASISWA & FRESH GRADUATES"
p_hm2.font.size = Pt(10.5)
p_hm2.font.bold = True
p_hm2.font.color.rgb = C_TEXT_WHITE
p_hm2.alignment = PP_ALIGN.CENTER

tb_m2 = slide12.shapes.add_textbox(Inches(1.1), Inches(2.35), Inches(5.133), Inches(4.2))
tf_m2 = tb_m2.text_frame
tf_m2.word_wrap = True
tf_m2.margin_left = tf_m2.margin_right = tf_m2.margin_top = tf_m2.margin_bottom = 0

m_benefits = [
    ("Penguasaan Narasi Portofolio Mandiri", "Mahasiswa tidak gagap saat diuji pewawancara karena mereka merumuskan sendiri poin pencapaiannya melalui dialog sokratik."),
    ("Efisiensi Tinggi Waktu Pembuatan Dokumen", "Reduksi waktu dari 3 jam riset manual menjadi < 5 menit dengan hasil terstruktur rapi dan 100% lolos screening ATS."),
    ("Kepercayaan Diri Menghadapi Wawancara", "Simulasi audio-visual STAR memberikan pengalaman langsung menjawab pertanyaan sulit dengan feedback objektif."),
    ("Panduan Peningkatan Keterampilan Riil", "Mengetahui secara tepat kursus atau sertifikasi apa yang harus diambil untuk menutup gap keterampilan karirnya.")
]

for t, d in m_benefits:
    p_t = tf_m2.add_paragraph() if tf_m2.paragraphs[0].text else tf_m2.paragraphs[0]
    p_t.text = f"✓ {t}"
    p_t.font.size = Pt(11)
    p_t.font.bold = True
    p_t.font.color.rgb = C_NAVY_DARK
    
    p_d = tf_m2.add_paragraph()
    p_d.text = f"   {d}"
    p_d.font.size = Pt(9.5)
    p_d.font.color.rgb = C_TEXT_MUTED

card_u = create_card(slide12, 6.8, 1.6, 5.733, 5.15, bg_color=C_CARD_BG, border_color=C_CYAN, border_width=1.5)
header_u = slide12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(1.8), Inches(5.333), Inches(0.45))
set_shape_fill(header_u, C_CYAN)
p_hu2 = header_u.text_frame.paragraphs[0]
p_hu2.text = "2. MANFAAT BAGI PERGURUAN TINGGI & CDC"
p_hu2.font.size = Pt(10.5)
p_hu2.font.bold = True
p_hu2.font.color.rgb = C_TEXT_WHITE
p_hu2.alignment = PP_ALIGN.CENTER

tb_u2 = slide12.shapes.add_textbox(Inches(7.1), Inches(2.35), Inches(5.133), Inches(4.2))
tf_u2 = tb_u2.text_frame
tf_u2.word_wrap = True
tf_u2.margin_left = tf_u2.margin_right = tf_u2.margin_top = tf_u2.margin_bottom = 0

u_benefits = [
    ("Peningkatan Indeks Kinerja Utama (IKU 1)", "Mendukung capaian IKU Perguruan Tinggi dalam mencetak lulusan yang langsung terserap kerja dengan gaji layak."),
    ("Data Analitik Kesenjangan Kurikulum", "Memberikan masukan berbasis data riil kepada ketua program studi mengenai kebutuhan kompetensi industri terkini."),
    ("Skalabilitas Bimbingan Karir Massal", "Menjadi co-pilot bagi konselor karir kampus untuk melayani ribuan mahasiswa secara simultan tanpa antrean."),
    ("Reputasi Kampus Berdaya Saing Digital", "Membuktikan kesiapan institusi dalam mengintegrasikan AI secara etis dan bertanggung jawab dalam proses akademik.")
]

for t, d in u_benefits:
    p_t = tf_u2.add_paragraph() if tf_u2.paragraphs[0].text else tf_u2.paragraphs[0]
    p_t.text = f"✓ {t}"
    p_t.font.size = Pt(11)
    p_t.font.bold = True
    p_t.font.color.rgb = C_NAVY_DARK
    
    p_d = tf_u2.add_paragraph()
    p_d.text = f"   {d}"
    p_d.font.size = Pt(9.5)
    p_d.font.color.rgb = C_TEXT_MUTED

add_footer(slide12)


# =============================================================
# SLIDE 13: PENILAIAN KELAYAKAN SISTEM (FEASIBILITY STUDY)
# =============================================================
slide13 = prs.slides.add_slide(blank_layout)
add_header(slide13, "Penilaian Kelayakan Sistem (Feasibility Analysis)", 
           subtitle="Evaluasi Menyeluruh Terhadap Aspek Operasional, Teknis, Waktu, dan Kelayakan Pedagogis-Ekonomi")

feasibilities = [
    {
        "num": "01",
        "title": "Kelayakan Operasional",
        "status": "SANGAT TINGGI (100% READY)",
        "desc": "Aplikasi berbasis Progressive Web App (PWA) berjalan lancar di smartphone Android, iOS, maupun laptop tanpa perlu instalasi rumit dari app store.",
        "color": C_TEAL
    },
    {
        "num": "02",
        "title": "Kelayakan Teknis",
        "status": "SANGAT TINGGI (15/15 ROUTES PASS)",
        "desc": "Teknologi modern Next.js 16, TypeScript, Golang microservice, dan Tailwind CSS memastikan arsitektur sangat modular, cepat, dan aman.",
        "color": C_CYAN
    },
    {
        "num": "03",
        "title": "Kelayakan Waktu",
        "status": "TERCAPAI 100% ON SCHEDULE",
        "desc": "Seluruh fase rekayasa dari perancangan kognitif, coding, security hardening, hingga uji coba 42 black-box test selesai tepat waktu sesuai deadline kompetisi.",
        "color": C_GOLD
    },
    {
        "num": "04",
        "title": "Kelayakan Pedagogis & Ekonomi",
        "status": "INOVATIF & SUSTAINABLE",
        "desc": "Menjadi solusi pertama yang secara tegas melarang otomatisasi buta AI melalui Explicit Approval Gate, dengan biaya operasional hosting $0 via Vercel.",
        "color": RGBColor(16, 185, 129)
    }
]

for idx, f in enumerate(feasibilities):
    col = idx % 2
    row = idx // 2
    left = 0.8 + col * 6.0
    top = 1.6 + row * 2.6
    
    card = create_card(slide13, left, top, 5.733, 2.4, bg_color=C_CARD_BG, border_color=C_CARD_BORDER)
    
    pill = slide13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left + 0.2), Inches(top + 0.2), Inches(0.65), Inches(0.4))
    set_shape_fill(pill, f["color"])
    p_num = pill.text_frame.paragraphs[0]
    p_num.text = f["num"]
    p_num.font.size = Pt(11)
    p_num.font.bold = True
    p_num.font.color.rgb = C_TEXT_WHITE
    p_num.alignment = PP_ALIGN.CENTER
    
    tb = slide13.shapes.add_textbox(Inches(left + 0.95), Inches(top + 0.15), Inches(4.6), Inches(2.1))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    
    p_t = tf.paragraphs[0]
    p_t.text = f["title"]
    p_t.font.size = Pt(13)
    p_t.font.bold = True
    p_t.font.color.rgb = C_NAVY_DARK
    
    p_st = tf.add_paragraph()
    p_st.text = f"Status: {f['status']}"
    p_st.font.size = Pt(9.5)
    p_st.font.bold = True
    p_st.font.color.rgb = f["color"]
    
    p_d = tf.add_paragraph()
    p_d.text = f"   {f['desc']}"
    p_d.font.size = Pt(10)
    p_d.font.color.rgb = C_TEXT_MUTED

add_footer(slide13)


# =============================================================
# SLIDE 14: TAMPILAN PROJEK & BUKTI DEPLOYMENT LIVE
# =============================================================
slide14 = prs.slides.add_slide(blank_layout)
add_header(slide14, "Tampilan Projek & Bukti Deployment Produksi", 
           subtitle="Aplikasi Bekerja Nyata (Working Prototype), Teruji 100%, dan Sudah Di-deploy Global")

dep_card = create_card(slide14, 0.8, 1.55, 11.733, 1.25, bg_color=C_NAVY_DARK, border_color=C_TEAL, border_width=1.5)
tf_dc = dep_card.text_frame
tf_dc.word_wrap = True

p_d1 = tf_dc.paragraphs[0]
r_d1a = p_d1.add_run()
r_d1a.text = "STATUS DEPLOYMENT: "
r_d1a.font.size = Pt(10)
r_d1a.font.bold = True
r_d1a.font.color.rgb = C_GOLD

r_d1b = p_d1.add_run()
r_d1b.text = "LIVE PRODUCTION READY (15/15 ROUTES PASS, 100% PWA COMPLIANT)"
r_d1b.font.size = Pt(10)
r_d1b.font.bold = True
r_d1b.font.color.rgb = C_TEAL

p_d2 = tf_dc.add_paragraph()
r_url_lbl = p_d2.add_run()
r_url_lbl.text = "🌐 Live URL Vercel : "
r_url_lbl.font.size = Pt(12)
r_url_lbl.font.bold = True
r_url_lbl.font.color.rgb = C_TEXT_WHITE

r_url_val = p_d2.add_run()
r_url_val.text = "https://career-student-ai.vercel.app"
r_url_val.font.size = Pt(13)
r_url_val.font.bold = True
r_url_val.font.color.rgb = C_CYAN

p_d3 = tf_dc.add_paragraph()
r_git_lbl = p_d3.add_run()
r_git_lbl.text = "📁 GitHub Repositori : "
r_git_lbl.font.size = Pt(11)
r_git_lbl.font.bold = True
r_git_lbl.font.color.rgb = C_TEXT_WHITE

r_git_val = p_d3.add_run()
r_git_val.text = "https://github.com/rikiardiansyah-web/CareerStudentAI"
r_git_val.font.size = Pt(11)
r_git_val.font.color.rgb = RGBColor(203, 213, 225)

card_ss1 = create_card(slide14, 0.8, 2.95, 5.733, 3.8, bg_color=C_CARD_BG, border_color=C_CARD_BORDER)
if os.path.exists(SS1_PATH):
    slide14.shapes.add_picture(SS1_PATH, Inches(0.95), Inches(3.05), width=Inches(5.433))

tb_c1 = slide14.shapes.add_textbox(Inches(0.95), Inches(6.15), Inches(5.433), Inches(0.5))
tf_c1 = tb_c1.text_frame
tf_c1.word_wrap = True
tf_c1.margin_left = tf_c1.margin_right = tf_c1.margin_top = tf_c1.margin_bottom = 0
p_c1 = tf_c1.paragraphs[0]
p_c1.text = "Modul 1: ATS Fast CV Builder dengan Explicit Approval Gate & Socratic Prompting"
p_c1.font.size = Pt(8.5)
p_c1.font.bold = True
p_c1.font.color.rgb = C_NAVY_MED
p_c1.alignment = PP_ALIGN.CENTER

card_ss2 = create_card(slide14, 6.8, 2.95, 5.733, 3.8, bg_color=C_CARD_BG, border_color=C_CARD_BORDER)
if os.path.exists(SS2_PATH):
    slide14.shapes.add_picture(SS2_PATH, Inches(6.95), Inches(3.05), width=Inches(5.433))

tb_c2 = slide14.shapes.add_textbox(Inches(6.95), Inches(6.15), Inches(5.433), Inches(0.5))
tf_c2 = tb_c2.text_frame
tf_c2.word_wrap = True
tf_c2.margin_left = tf_c2.margin_right = tf_c2.margin_top = tf_c2.margin_bottom = 0
p_c2 = tf_c2.paragraphs[0]
p_c2.text = "Modul 2 & 3: Skill-Gap Matrix (Taksonomi WEF 2023) & Mock Interview STAR Framework"
p_c2.font.size = Pt(8.5)
p_c2.font.bold = True
p_c2.font.color.rgb = C_NAVY_MED
p_c2.alignment = PP_ALIGN.CENTER

add_footer(slide14)


# =============================================================
# SLIDE 15: PENUTUP / TERIMA KASIH (Dark Navy Theme)
# =============================================================
slide15 = prs.slides.add_slide(blank_layout)
bg15 = slide15.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
set_shape_fill(bg15, C_NAVY_DARK)

bar15 = slide15.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
set_shape_fill(bar15, C_TEAL)

if os.path.exists(LOGO_PATH):
    logo_card15 = slide15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.8), Inches(1.0), Inches(1.733), Inches(1.733))
    set_shape_fill(logo_card15, RGBColor(255, 255, 255), C_TEAL, border_width=2)
    slide15.shapes.add_picture(LOGO_PATH, Inches(5.95), Inches(1.15), width=Inches(1.433))

close_box = slide15.shapes.add_textbox(Inches(1.5), Inches(2.9), Inches(10.333), Inches(3.6))
tf_cl = close_box.text_frame
tf_cl.word_wrap = True
tf_cl.margin_left = tf_cl.margin_right = tf_cl.margin_top = tf_cl.margin_bottom = 0

p_c_thx = tf_cl.paragraphs[0]
p_c_thx.text = "TERIMA KASIH"
p_c_thx.font.size = Pt(40)
p_c_thx.font.bold = True
p_c_thx.font.color.rgb = C_TEXT_WHITE
p_c_thx.alignment = PP_ALIGN.CENTER

p_c_tag = tf_cl.add_paragraph()
p_c_tag.text = "\"Menjembatani Nalar Kritis Mahasiswa di Era Disrupsi Kecerdasan Buatan\""
p_c_tag.font.size = Pt(14)
p_c_tag.font.italic = True
p_c_tag.font.color.rgb = C_TEAL
p_c_tag.alignment = PP_ALIGN.CENTER

p_spc = tf_cl.add_paragraph()
p_spc.text = ""
p_spc.font.size = Pt(8)

p_c_team = tf_cl.add_paragraph()
p_c_team.text = "TIM AQUAMAN TEKNOKRAT"
p_c_team.font.size = Pt(20)
p_c_team.font.bold = True
p_c_team.font.color.rgb = C_GOLD
p_c_team.alignment = PP_ALIGN.CENTER

p_c_names = tf_cl.add_paragraph()
p_c_names.text = "Riki Ardiansyah (Ketua)  •  Jovan Akbar Setiodono  •  Fiqi Rahman Al Afkhoni"
p_c_names.font.size = Pt(12)
p_c_names.font.color.rgb = RGBColor(226, 232, 240)
p_c_names.alignment = PP_ALIGN.CENTER

p_c_camp = tf_cl.add_paragraph()
p_c_camp.text = "UNIVERSITAS TEKNOKRAT INDONESIA — KAMPUSNYA SANG JUARA"
p_c_camp.font.size = Pt(13)
p_c_camp.font.bold = True
p_c_camp.font.color.rgb = C_TEXT_WHITE
p_c_camp.alignment = PP_ALIGN.CENTER

cta_btn = slide15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.5), Inches(5.8), Inches(4.333), Inches(0.55))
set_shape_fill(cta_btn, C_TEAL, border_width=0)
p_cta = cta_btn.text_frame.paragraphs[0]
p_cta.text = "Akses Prototipe: career-student-ai.vercel.app"
p_cta.font.size = Pt(11)
p_cta.font.bold = True
p_cta.font.color.rgb = C_TEXT_WHITE
p_cta.alignment = PP_ALIGN.CENTER

add_footer(slide15, is_dark=True)

prs.save(OUTPUT_SCRATCH)
prs.save(OUTPUT_BRAIN)
print("Updated Pitch Deck saved successfully with Golang microservices included!")
